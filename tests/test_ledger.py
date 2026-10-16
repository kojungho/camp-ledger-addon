import importlib.util, io, json, os, tempfile, unittest, zipfile, sys
from unittest.mock import patch
from pathlib import Path
from datetime import datetime
from PIL import Image
from openpyxl import load_workbook

sandbox=tempfile.TemporaryDirectory()
os.environ['CAMP_DATA']=sandbox.name
os.environ.pop('HA_INGRESS',None)
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'camp_ledger'))
spec=importlib.util.spec_from_file_location('ledger',Path(__file__).resolve().parents[1]/'camp_ledger/app.py')
ledger=importlib.util.module_from_spec(spec);spec.loader.exec_module(ledger)

class LedgerTest(unittest.TestCase):
    def setUp(self):
        self.client=ledger.app.test_client()
        self.headers={'X-Camp-Request':'1'}
        with ledger.conn() as c:
            c.execute('DELETE FROM items')
            c.execute('DELETE FROM imports')
            c.execute('DELETE FROM devices')
            c.execute('UPDATE settings SET body=? WHERE id=1',(json.dumps(ledger.DEFAULT),))
    def test_calendar_choice_settings_roundtrip_and_invalid_size(self):
        settings=self.client.get('/api/state').json['settings'];self.assertEqual(settings['display']['calendarChoice'],{'width':220,'height':210,'font':14})
        settings['display']['calendarChoice']={'width':180,'height':160,'font':12};self.assertEqual(self.client.put('/api/settings',json=settings,headers=self.headers).status_code,200);self.assertEqual(self.client.get('/api/state').json['settings']['display']['calendarChoice'],settings['display']['calendarChoice'])
        for key,value in [('width',159),('height',481),('font',9),('width',True)]:
            bad=json.loads(json.dumps(settings));bad['display']['calendarChoice'][key]=value;self.assertEqual(self.client.put('/api/settings',json=bad,headers=self.headers).status_code,400)
    def image(self):
        b=io.BytesIO();Image.new('RGB',(32,24),'green').save(b,'PNG');b.seek(0);return b
    def item(self,**extra):
        return dict(category='텐트',name='캠핑 텐트',purchase_date='2026-10-01',method='온라인',buyer='나',purchase_price=300000,sale_price=150000,status='owned',sale_date='',disposal_date='',notes='영수증 보관',**extra)
    def save(self,item,ident=None,photos=0,deletes=None):
        data={'item':json.dumps(item),'delete_attachments':json.dumps(deletes or [])}
        if photos:data['photo_kinds']=json.dumps(['product' if n%2==0 else 'receipt' for n in range(photos)])
        if photos:data['photos']=[(self.image(),f'사진{n}.png') for n in range(photos)]
        return self.client.open('/api/items'+('/'+ident if ident else ''),method='PUT' if ident else 'POST',data=data,headers=self.headers)
    def state(self):return self.client.get('/api/state').get_json()
    def test_staged_records_undo_and_validation(self):
        r=self.save(self.item(purchase_qty=6,current_qty=6,sold_qty=0)); ident=r.json['id']
        old=self.state()['items'][0]
        edit={**old,'purchase_qty':8,'current_qty':7,'purchase_price':400000,'operations':[{'kind':'restock','quantity':2,'cost':100000,'date':'2026-10-14','purchase_type':'선물','notes':'추가'},{'kind':'consume','quantity':1,'date':'2026-10-14','notes':'사용'}]}
        self.assertEqual(self.save(edit,ident).status_code,200)
        item=self.state()['items'][0]; self.assertEqual(len(item['purchases']),2); self.assertEqual(len(item['consumptions']),1); self.assertEqual(item['current_value'],350000)
        bad={**item,'current_qty':9}; r=self.save(bad,ident); self.assertEqual(r.status_code,400);self.assertIn('current_qty',r.json['fields']);self.assertEqual(self.state()['items'][0]['current_qty'],7)
        undo={**item,'purchase_qty':6,'current_qty':6,'purchase_price':300000,'operations':[{'kind':'undo_consume','id':item['consumptions'][0]['id']},{'kind':'undo_purchase','id':item['purchases'][1]['id']}]}
        self.assertEqual(self.save(undo,ident).status_code,200)
        item=self.state()['items'][0];self.assertEqual(len(item['purchases']),1);self.assertEqual(item['consumptions'],[]);self.assertEqual(item['current_qty'],6)
        s=self.state()['settings'];s['purchase_types']=['선물','중고'];s['display']['cellPadding']=0;s['display']['columns']['name'].update(width=24,auto=True);s['display']['columnOrder']=list(reversed(ledger.DISPLAY_LABELS));s['display']['visibleColumns']=['name','category']
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200);self.assertEqual(self.state()['settings']['purchase_types'],['선물','중고'])

    def test_live_revision_and_setting_conflicts(self):
        revision=self.client.get('/api/revision').json['revision'];before=self.state()
        self.assertEqual(before['revision'],revision)
        self.assertEqual(self.client.get('/api/revision').json['revision'],revision)
        ident=self.save(self.item(),photos=2).json['id'];created=self.state()
        self.assertGreater(created['revision'],revision);self.assertEqual(created['settings_version'],before['settings_version'])
        invalid={**created['items'][0],'purchase_qty':0};self.assertEqual(self.save(invalid,ident).status_code,400)
        self.assertEqual(self.state()['revision'],created['revision'])
        s={**created['settings'],'expected_version':created['settings_version']};s['display']['cellPadding']=5
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
        newer=self.state();self.assertGreater(newer['revision'],created['revision']);self.assertNotEqual(newer['settings_version'],created['settings_version'])
        s['display']['cellPadding']=8
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,409)
        self.assertEqual(self.state()['settings']['display']['cellPadding'],5)
        with ledger.conn() as c:
            c.execute('BEGIN IMMEDIATE');c.execute('DELETE FROM items WHERE id=?',(ident,));c.rollback()
        self.assertEqual(self.state()['revision'],newer['revision'])
        stream=self.client.get('/api/events',buffered=False);parts=iter(stream.response)
        self.assertIn(b'retry:',next(parts));self.assertIn(b'event: change',next(parts))
        self.client.post('/api/delete',json={'ids':[ident]},headers=self.headers)
        self.assertIn(b'event: change',next(parts));stream.close()
        self.assertGreater(self.state()['revision'],newer['revision'])

    def test_dollar_purchases_and_mobile_settings(self):
        rate={'exchange_rate':1333.09,'exchange_date':'2026-10-12'}
        with patch.object(ledger,'get_rate',return_value=rate):
            ident=self.save({**self.item(purchase_currency='USD',foreign_amount=100.25,purchase_qty=2,current_qty=1,shipping_fee=3000),'purchase_date':'2026-10-12'}).json['id']
            item=self.state()['items'][0];self.assertEqual(item['purchase_price'],133642);self.assertEqual(item['current_value'],68321)
            self.assertEqual(item['purchases'][0]['foreign_amount'],100.25)
            edit={**item,'purchase_qty':3,'current_qty':2,'purchase_price':147639,'shipping_fee':4000,'operations':[{'kind':'restock','quantity':1,'cost':13997,'shipping_fee':1000,'date':'2026-10-12','purchase_currency':'USD','foreign_amount':10.50}]}
            self.assertEqual(self.save(edit,ident).status_code,200)
            item=self.state()['items'][0];self.assertEqual(item['purchase_price'],147639);self.assertEqual(item['purchases'][1]['cost'],13997)
            self.assertEqual(self.save({**item,'foreign_amount':1.001},ident).status_code,400)
        wb=load_workbook(io.BytesIO(self.client.get('/api/export').data));headers=[c.value for c in wb['전체 상품'][1]];date_cell=wb['전체 상품'].cell(2,headers.index('구입일자')+1)
        self.assertEqual(date_cell.number_format,'yy.mm.dd');self.assertIn('달러 구매금액',[c.value for c in wb['구매 내역'][1]])
        backup=self.client.get('/api/backup').data
        with patch.object(ledger,'get_rate',side_effect=ValueError('환율 연결 실패')):
            self.assertEqual(self.client.post('/api/restore',data={'backup':(io.BytesIO(backup),'backup.zip')},headers=self.headers).status_code,200)
            saved=self.state()['items'][0];self.assertEqual(saved['purchase_price'],147639)
            self.assertEqual(self.save(saved,saved['id']).status_code,400)
        s=self.state()['settings'];m=s['display']['mobile'];self.assertEqual(m['headerHeight'],0);self.assertEqual(m['headerPadding'],0)
        m.update(labelGap=1,headerFont=9);m['columns']['brand']={'fontWeight':700,'padding':0}
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
        self.assertEqual(self.state()['settings']['display']['headerHeight'],48)
        m['labelGap']=-1;self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,400)

    def test_fine_display_preferences(self):
        s=self.state()['settings'];d=s['display'];self.assertEqual(d['headerFont'],12);self.assertEqual(d['bodyWeight'],400)
        d.update(headerFont=18,headerPadding=0,headerHeight=28,bodyWeight=600,cellVerticalPadding=0)
        d['columns']['brand'].update(fontWeight=700,headerFont=20,headerWeight=800,padding=0,verticalPadding=2,rowHeight=26,headerPadding=1,headerHeight=24)
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
        saved=self.state()['settings'];self.assertEqual(saved['display']['columns']['brand']['fontWeight'],700)
        saved['display']['columns']['brand']['fontWeight']=1000
        self.assertEqual(self.client.put('/api/settings',json=saved,headers=self.headers).status_code,400)
        saved['display']['columns']['brand']['fontWeight']=None
        self.assertEqual(self.client.put('/api/settings',json=saved,headers=self.headers).status_code,200)

    def test_shipping_cost_lifecycle(self):
        ident=self.save(self.item(purchase_qty=6,current_qty=5,shipping_fee=6000)).json['id']
        item=self.state()['items'][0]
        self.assertEqual((item['total_purchase_cost'],item['avg_unit_price'],item['current_value'],item['deducted_value']),(306000,51000,255000,51000))
        edit={**item,'purchase_qty':8,'current_qty':7,'purchase_price':400000,'shipping_fee':10000,'operations':[{'kind':'restock','quantity':2,'cost':100000,'shipping_fee':4000,'date':'2026-10-14'}]}
        self.assertEqual(self.save(edit,ident).status_code,200)
        item=self.state()['items'][0];self.assertEqual(item['total_purchase_cost'],410000);self.assertEqual([p['shipping_fee'] for p in item['purchases']],[6000,4000])
        bad={**item,'shipping_fee':-1};self.assertEqual(self.save(bad,ident).status_code,400)
        bad={**item,'shipping_fee':3000};self.assertEqual(self.save(bad,ident).status_code,400)
        wb=load_workbook(io.BytesIO(self.client.get('/api/export').data));rows=list(wb['전체 상품'].values);record=dict(zip(rows[0],rows[1]));self.assertEqual(record['배송비'],10000);self.assertEqual(record['배송비 포함 구입금액'],410000)
        self.assertIn('배송비',list(wb['구매 내역'].values)[0]);self.assertEqual(list(wb['금액 집계'].values)[4][2],410000)
        backup=self.client.get('/api/backup').data
        self.assertEqual(self.client.post('/api/restore',data={'backup':(io.BytesIO(backup),'backup.zip')},headers=self.headers).status_code,200)
        item=self.state()['items'][0];self.assertEqual(item['shipping_fee'],10000)
        undo={**item,'purchase_qty':6,'current_qty':5,'purchase_price':300000,'shipping_fee':6000,'operations':[{'kind':'undo_purchase','id':item['purchases'][1]['id']}]}
        self.assertEqual(self.save(undo,ident).status_code,200);self.assertEqual(self.state()['items'][0]['total_purchase_cost'],306000)
        r=self.client.post('/api/items/'+ident+'/restock',json={'updated_at':self.state()['items'][0]['updated_at'],'quantity':1,'cost':20000,'shipping_fee':3000},headers=self.headers)
        self.assertEqual(r.status_code,200);self.assertEqual(self.state()['items'][0]['total_purchase_cost'],329000)
        legacy=self.item();legacy.update(id='legacy',history=[])
        self.assertEqual(ledger.stock_view(legacy)['shipping_fee'],0)

    def test_product_and_receipt_delete_then_readd_same_name(self):
        ident=self.save(self.item(),photos=2).json['id']
        for cycle in range(2):
            old=self.state()['items'][0];deleted=[a['id'] for a in old['attachments']]
            self.assertEqual(self.save(old,ident,deletes=deleted).status_code,200)
            cleared=self.state()['items'][0];self.assertEqual(cleared['attachments'],[])
            self.assertEqual(self.save(cleared,ident,photos=2).status_code,200)
            restored=self.state()['items'][0]
            self.assertEqual([a['kind'] for a in restored['attachments']],['product','receipt'])
            self.assertEqual([a['name'] for a in restored['attachments']],['사진0.png','사진1.png'])
            self.assertEqual((restored['product_photo_count'],restored['receipt_photo_count']),(1,1))
            for attachment in restored['attachments']:self.assertEqual(self.client.get('/api/photos/'+attachment['id']).status_code,200)
            for aid in deleted:self.assertEqual(self.client.get('/api/photos/'+aid).status_code,404)

    def test_brand_link_and_mobile_preferences(self):
        ident=self.save(self.item(brand='네이처하이크',purchase_url='https://example.com/item?ref=camp')).json['id']
        saved=self.state()['items'][0];self.assertEqual(saved['brand'],'네이처하이크');self.assertEqual(saved['purchase_url'],'https://example.com/item?ref=camp')
        bad={**saved,'purchase_url':'javascript:alert(1)'};r=self.save(bad,ident);self.assertEqual(r.status_code,400);self.assertEqual(r.json['fields'],['purchase_url'])
        wb=load_workbook(io.BytesIO(self.client.get('/api/export').data));self.assertIn('브랜드',[c.value for c in wb['전체 상품'][1]]);self.assertIn('구매 사이트 링크',[c.value for c in wb['전체 상품'][1]])
        s=self.state()['settings'];self.assertFalse(s['display']['rememberDialogScroll']);self.assertEqual(s['display']['mobileColumns'],2)
        s['display'].update(mobileColumns=8,rememberDialogScroll=True);self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
        s['display']['mobileColumns']=9;self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,400)
        self.assertEqual(self.state()['settings']['display']['mobileColumns'],8)

    def test_full_lifecycle_and_export(self):
        r=self.save(self.item(),photos=12);self.assertEqual(r.status_code,200)
        first=self.state()['items'][0];self.assertEqual(len(first['attachments']),12);self.assertEqual(sum(a['kind']=='receipt' for a in first['attachments']),6)
        changed={**first,'status':'sold','sale_date':'2026-10-13','sale_price':180000,'name':'=SUM(1,2)'}
        r=self.save(changed,first['id'],photos=1,deletes=[first['attachments'][0]['id']]);self.assertEqual(r.status_code,200)
        sold=self.state()['items'][0];self.assertEqual(len(sold['attachments']),12);self.assertEqual(len(sold['history']),2)
        self.assertEqual(self.client.get('/api/photos/'+first['attachments'][0]['id']).status_code,404)
        x=self.client.get('/api/export');self.assertEqual(x.status_code,200)
        wb=load_workbook(io.BytesIO(x.data));self.assertEqual(wb.sheetnames,['전체 상품','변경 이력','첨부 사진','설정','금액 집계','구매 내역','사용·소진 내역'])
        self.assertEqual(wb['전체 상품']['C2'].value,"'=SUM(1,2)")
        self.assertEqual(len(wb['첨부 사진']._images),12)
        self.assertEqual(wb['금액 집계']['D3'].value,180000)
        self.assertEqual(wb['금액 집계']['E5'].value,120000)
        disposed={**sold,'status':'disposed','disposal_date':'2026-10-13'}
        self.assertEqual(self.save(disposed,sold['id']).status_code,200)
        wb=load_workbook(io.BytesIO(self.client.get('/api/export').data));self.assertEqual(wb['금액 집계']['D5'].value,0)
        backup=self.client.get('/api/backup').data
        r=self.client.post('/api/delete',json={'ids':[first['id']]},headers=self.headers);self.assertEqual(r.status_code,200);self.assertEqual(self.state()['items'],[])
        with ledger.conn() as c:self.assertEqual(c.execute('SELECT COUNT(*) FROM attachments').fetchone()[0],0)
        r=self.client.post('/api/restore',data={'backup':(io.BytesIO(backup),'backup.zip')},headers=self.headers)
        self.assertEqual(r.status_code,200);restored=self.state()['items'][0];self.assertEqual(restored['status'],'disposed');self.assertEqual(len(restored['attachments']),12)
        self.assertEqual(self.client.get('/api/photos/'+restored['attachments'][0]['id']).status_code,200)
    def test_validation_conflict_and_atomic_restore(self):
        bad={**self.item(),'purchase_price':-1};self.assertEqual(self.save(bad).status_code,400)
        r=self.save(self.item());self.assertEqual(r.status_code,200);first=self.state()['items'][0]
        self.assertEqual(self.save({**first,'notes':'수정'},first['id']).status_code,200)
        self.assertEqual(self.save(first,first['id']).status_code,409)
        r=self.client.post('/api/items',data={'item':json.dumps(self.item()),'photos':(io.BytesIO(b'fake'),'fake.png')},headers=self.headers);self.assertEqual(r.status_code,400)
        self.assertEqual(len(self.state()['items']),1)
        z=io.BytesIO()
        with zipfile.ZipFile(z,'w') as archive:archive.writestr('ledger.json',json.dumps({'version':1,'settings':ledger.DEFAULT,'items':[{'bad':1}]}))
        z.seek(0);self.assertEqual(self.client.post('/api/restore',data={'backup':(z,'bad.zip')},headers=self.headers).status_code,400)
        self.assertEqual(len(self.state()['items']),1)
    def test_settings_persistence_and_ingress(self):
        s=self.state()['settings'];s['categories'].remove('텐트');s['buyers'].append('캠퍼');s['display']['columns']['name']['font']=22
        self.save(self.item());r=self.client.put('/api/settings',json=s,headers=self.headers);self.assertEqual(r.status_code,200)
        self.assertEqual(self.state()['settings']['display']['columns']['name']['font'],22)
        self.assertEqual(self.state()['items'][0]['category'],'텐트')
        self.assertEqual(self.client.post('/api/delete',json={'ids':[]}).status_code,400)
        os.environ['HA_INGRESS']='1'
        try:
            self.assertEqual(self.client.get('/api/state').status_code,403)
            self.assertEqual(self.client.get('/api/state',environ_overrides={'REMOTE_ADDR':'172.30.32.2','SERVER_PORT':'8099'}).status_code,200)
        finally:os.environ.pop('HA_INGRESS',None)

    def excel_upload(self,rows=None,headers=None):
        workbook=load_workbook(io.BytesIO(self.client.get('/api/import/template').data))
        ws=workbook['용품 입력']
        if headers:
            for n,header in enumerate(headers,1):ws.cell(1,n,header)
        for n,row in enumerate(rows or [],2):
            for col,value in enumerate(row,1):
                if value is not None:ws.cell(n,col,value)
        buf=io.BytesIO();workbook.save(buf);buf.seek(0)
        return self.client.post('/api/import/preview',data={'excel':(buf,'새용품.xlsx')},headers=self.headers)

    def test_excel_template_and_append_only_import(self):
        original=self.save(self.item(),photos=2).get_json()['id']
        before=self.state()['items'][0]
        workbook=load_workbook(io.BytesIO(self.client.get('/api/import/template').data))
        self.assertEqual(workbook['용품 입력']['B1'].value,'상품명')
        self.assertIsNone(workbook['용품 입력']['B2'].value)
        self.assertEqual(workbook['목록'].sheet_state,'hidden')
        self.assertEqual(len(workbook['용품 입력'].data_validations.dataValidation),12)
        self.assertEqual(self.excel_upload().status_code,400)
        rows=[['텐트','캠핑 텐트',datetime(2026,10,12),'온라인','나',310000,None,None,None,None,'동일 이름도 추가'],
              ['조명','랜턴','2026.10.13','직거래','캠퍼','₩35,000원','20,000','판매 완료','2026-10-13',None,None],
              [None,None,None,None,None,None,None,None,None,None,None],
              ['부품','폐기 부품','20261001',None,None,1000,0,'폐기',None,'2026/10/13','파손']]
        response=self.excel_upload(rows);self.assertEqual(response.status_code,200,response.get_json())
        result=response.get_json();self.assertTrue(result['ok']);self.assertEqual(result['count'],3)
        self.assertEqual(result['purchase_total'],346000);self.assertEqual(result['sale_total'],20000)
        self.assertEqual(len(self.state()['items']),1)
        commit=self.client.post('/api/import/commit',json={'token':result['token']},headers=self.headers)
        self.assertEqual(commit.status_code,200);self.assertEqual(commit.get_json()['count'],3)
        state=self.state();self.assertEqual(len(state['items']),4)
        retained=next(i for i in state['items'] if i['id']==original);self.assertEqual(retained,before)
        self.assertEqual(sum(i['name']=='캠핑 텐트' for i in state['items']),2)
        self.assertIn('조명',state['settings']['categories']);self.assertIn('직거래',state['settings']['methods']);self.assertIn('캠퍼',state['settings']['buyers'])
        imported=next(i for i in state['items'] if i['name']=='랜턴')
        self.assertEqual(imported['status'],'sold');self.assertEqual(imported['history'][0]['changes'][0]['field'],'엑셀 등록')
        repeat=self.client.post('/api/import/commit',json={'token':result['token']},headers=self.headers)
        self.assertEqual(repeat.get_json(),commit.get_json());self.assertEqual(len(self.state()['items']),4)
        new_preview=self.excel_upload(rows).get_json()
        self.assertEqual(self.client.post('/api/import/commit',json={'token':new_preview['token']},headers=self.headers).status_code,200)
        self.assertEqual(len(self.state()['items']),7)

    def test_excel_errors_do_not_write_partial_records(self):
        self.save(self.item(),photos=1);before=self.state()
        result=self.excel_upload([['의자','정상 의자',None,None,None,50000],
            ['의자','오류 의자','2026-02-30',None,None,-1,3.5,'미정'],
            ['텐트',None,None,None,None,100],
            ['텐트','=SUM(1,2)',None,None,None,100]]).get_json()
        self.assertFalse(result['ok']);self.assertNotIn('token',result);self.assertGreaterEqual(result['error_count'],6)
        self.assertEqual(self.state(),before)
        self.assertEqual(self.client.post('/api/import/commit',json={'token':'invalid'},headers=self.headers).status_code,400)
        bad=self.client.post('/api/import/preview',data={'excel':(io.BytesIO(b'fake'),'bad.xlsx')},headers=self.headers)
        self.assertEqual(bad.status_code,400)
        wrong=self.client.post('/api/import/preview',data={'excel':(io.BytesIO(b'fake'),'bad.xls')},headers=self.headers)
        self.assertEqual(wrong.status_code,400);self.assertEqual(self.state(),before)

    def test_export_can_be_appended_and_preview_expiry(self):
        self.save(self.item())
        exported=self.client.get('/api/export').data
        preview=self.client.post('/api/import/preview',data={'excel':(io.BytesIO(exported),'export.xlsx')},headers=self.headers).get_json()
        self.assertTrue(preview['ok']);self.assertEqual(preview['count'],1)
        with ledger.conn() as c:c.execute('UPDATE imports SET created=0 WHERE id=?',(preview['token'],))
        self.assertEqual(self.client.post('/api/import/commit',json={'token':preview['token']},headers=self.headers).status_code,400)
        self.assertEqual(len(self.state()['items']),1)

    def test_scroll_top_position_settings(self):
        settings=self.state()['settings']
        self.assertEqual(settings['display']['scrollTopPosition'],'bottom-right')
        for position in ('top-left','top-center','top-right','bottom-left','bottom-center','bottom-right'):
            settings['display']['scrollTopPosition']=position
            self.assertEqual(self.client.put('/api/settings',json=settings,headers=self.headers).status_code,200)
            self.assertEqual(self.state()['settings']['display']['scrollTopPosition'],position)
        settings['display']['scrollTopPosition']='middle'
        self.assertEqual(self.client.put('/api/settings',json=settings,headers=self.headers).status_code,400)
        self.assertEqual(self.state()['settings']['display']['scrollTopPosition'],'bottom-right')
        settings['display'].pop('scrollTopPosition')
        with ledger.conn() as c:c.execute('UPDATE settings SET body=? WHERE id=1',(json.dumps(settings),))
        self.assertEqual(self.state()['settings']['display']['scrollTopPosition'],'bottom-right')

    def test_photo_picker_settings(self):
        s=self.state()['settings']
        self.assertEqual(s['display']['photoPickerSource'],'files')
        for source in ('gallery','files'):
            s['display']['photoPickerSource']=source
            self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
            self.assertEqual(self.state()['settings']['display']['photoPickerSource'],source)
        s['display']['photoPickerSource']='camera'
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,400)
        self.assertEqual(self.state()['settings']['display']['photoPickerSource'],'files')

    def test_mobile_item_spacing_persistence_and_validation(self):
        s=self.state()['settings']
        pc=json.loads(json.dumps(s['display']['columns']))
        spacing={'width':20,'labelGap':0,'spaceTop':3,'spaceBottom':7,'spaceLeft':2,'spaceRight':9}
        s['display']['mobile']['columns']={'category':spacing,'name':{'spaceBottom':6}}
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
        saved=self.state()['settings']['display']
        self.assertEqual(saved['mobile']['columns']['category'],spacing)
        self.assertEqual(saved['columns'],pc)
        for key,value in [('width',19),('width',601),('labelGap',41),('spaceTop',-1),('spaceRight',61)]:
            invalid=json.loads(json.dumps(s));invalid['display']['mobile']['columns']['category'][key]=value
            self.assertEqual(self.client.put('/api/settings',json=invalid,headers=self.headers).status_code,400)
        self.assertEqual(self.state()['settings']['display']['mobile']['columns']['category'],spacing)

    def test_all_display_fields_and_legacy_settings(self):
        self.save(self.item())
        s=self.state()['settings'];s['display']['visibleColumns']=list(ledger.DISPLAY_LABELS)
        s['display']['visibleSummaries']=['purchase','owned'];s['display']['columns']['notes']={'font':18,'width':400}
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
        self.assertEqual(self.state()['settings']['display']['visibleColumns'],list(ledger.DISPLAY_LABELS))
        s['display']['visibleColumns']=['buyer','status'];s['display']['visibleSummaries']=[]
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,200)
        wb=load_workbook(io.BytesIO(self.client.get('/api/export').data))
        self.assertIn('상품명',[c.value for c in wb['전체 상품'][1]])
        self.assertEqual(len(self.state()['items']),1)
        s['display']['visibleColumns']=[]
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,400)
        s['display']['visibleColumns']=['unknown']
        self.assertEqual(self.client.put('/api/settings',json=s,headers=self.headers).status_code,400)
        legacy=json.loads(json.dumps(ledger.DEFAULT))
        for key in ('visibleColumns','visibleSummaries','showProductThumb'):legacy['display'].pop(key)
        legacy['display']['columns']={k:v for k,v in legacy['display']['columns'].items() if k in ('category','name','purchase_price','sale_price')}
        with ledger.conn() as c:c.execute('UPDATE settings SET body=? WHERE id=1',(json.dumps(legacy),))
        updated=self.state()['settings']['display']
        self.assertEqual(updated['visibleColumns'],['category','name','purchase_price','sale_price'])
        self.assertIn('notes',updated['columns']);self.assertTrue(updated['showProductThumb'])

    def test_quantity_weighted_average_and_used_restock(self):
        item={**self.item(),'purchase_price':60000,'sale_price':0,'purchase_qty':6,'current_qty':5,'sold_qty':0,'purchase_type':'새상품'}
        self.assertEqual(self.save(item).status_code,200)
        old=self.state()['items'][0]
        self.assertEqual((old['avg_unit_price'],old['current_value'],old['deducted_value'],old['disposed_qty']),(10000,50000,10000,1))
        r=self.client.post('/api/items/'+old['id']+'/restock',json={'quantity':6,'cost':90000,'date':'2026-10-13','notes':'추가 구매','purchase_type':'중고','updated_at':old['updated_at']},headers=self.headers)
        self.assertEqual(r.status_code,200)
        new=self.state()['items'][0]
        self.assertEqual((new['purchase_qty'],new['current_qty'],new['purchase_price']),(12,11,150000))
        self.assertEqual((new['avg_unit_price'],new['current_value'],new['deducted_value']),(12500,137500,12500))
        self.assertEqual(new['purchase_type'],'혼합');self.assertEqual(len(new['purchases']),2)
        self.assertEqual(self.save({**new,'notes':'수정'},new['id']).status_code,200)
        edited=self.state()['items'][0];self.assertEqual(len(edited['purchases']),2)
        self.assertEqual(edited['purchase_type'],'혼합')
        self.assertEqual(self.save({**edited,'current_qty':13},new['id']).status_code,400)
        wb=load_workbook(io.BytesIO(self.client.get('/api/export').data));self.assertEqual(wb['구매 내역'].max_row,3)
        backup=self.client.get('/api/backup').data
        self.assertEqual(self.client.post('/api/restore',data={'backup':(io.BytesIO(backup),'backup.zip')},headers=self.headers).status_code,200)
        self.assertEqual(self.state()['items'][0]['current_value'],137500)

    def test_shared_purchase_cost_remains_unconfirmed(self):
        self.save({**self.item(),'name':'공동구매 텐트','purchase_price':30000,'cost_unconfirmed':True,'shared_purchase_total':30000})
        self.save({**self.item(),'name':'공동구매 의자','purchase_price':0,'sale_price':75000,'status':'sold','purchase_qty':4,'current_qty':0,'sold_qty':4,'cost_unconfirmed':True,'shared_purchase_total':30000})
        state=self.state();self.assertEqual(sum(i['purchase_price'] for i in state['items']),30000)
        self.assertTrue(all(i['avg_unit_price'] is None for i in state['items']))
        self.assertTrue(all(i['current_value'] is None for i in state['items']))
        chair=next(i for i in state['items'] if i['name']=='공동구매 의자');self.assertEqual(chair['sold_qty'],4)

    def test_device_access_scope_and_revocation(self):
        self.save(self.item())
        os.environ['HA_INGRESS']='1'
        owner={'REMOTE_ADDR':'172.30.32.2','SERVER_PORT':'8099'}
        outside={'REMOTE_ADDR':'172.30.32.2','SERVER_PORT':'8098','camp.entrypoint':'external'}
        external=ledger.app.test_client(use_cookies=False)
        try:
            self.assertEqual(external.get('/api/state',environ_overrides=outside).status_code,401)
            self.assertEqual(external.get('/api/state',headers={'Host':'ha.badara.co.kr:8099'},environ_overrides={**outside,'SERVER_PORT':'8099'}).status_code,401)
            page=external.get('/',environ_overrides=outside)
            self.assertEqual(page.status_code,200);self.assertNotIn('캠핑 텐트',page.data.decode());page.close()
            cfg=self.client.put('/api/mobile/config',json={'enabled':True,'external_url':'https://ha.badara.co.kr/camp-ledger/'},headers=self.headers,environ_overrides=owner)
            self.assertEqual(cfg.status_code,200)
            pair=self.client.post('/api/mobile/devices',json={'name':'테스트 휴대폰'},headers=self.headers,environ_overrides=owner).get_json()
            self.assertGreater(len(pair['token']),32)
            with ledger.conn() as c:self.assertNotEqual(c.execute('SELECT token_hash FROM devices').fetchone()[0],pair['token'])
            cookie={**self.headers,'Cookie':'camp_device='+pair['token']}
            self.assertEqual(external.get('/api/state',headers=cookie,environ_overrides=outside).status_code,200)
            self.assertFalse(external.get('/api/state',headers=cookie,environ_overrides=outside).get_json()['mobile']['admin'])
            self.assertEqual(external.post('/api/mobile/devices',json={'name':'forbidden'},headers=cookie,environ_overrides=outside).status_code,403)
            self.assertEqual(external.put('/api/mobile/config',json={'enabled':False,'external_url':''},headers=cookie,environ_overrides=outside).status_code,403)
            s=external.get('/api/state',headers=cookie,environ_overrides=outside).get_json()['settings'];s['mobile']={'enabled':False,'external_url':'https://wrong.example/'}
            self.assertEqual(external.put('/api/settings',json=s,headers=cookie,environ_overrides=outside).status_code,200)
            self.assertTrue(self.client.get('/api/mobile/status',environ_overrides=owner).get_json()['enabled'])
            login=external.post('/api/mobile/session',json={'key':pair['token']},headers=self.headers,environ_overrides=outside)
            self.assertEqual(login.status_code,200);cookie_header=login.headers['Set-Cookie']
            self.assertIn('Secure',cookie_header);self.assertIn('HttpOnly',cookie_header);self.assertIn('Path=/camp-ledger/',cookie_header)
            self.assertEqual(self.client.post('/api/mobile/devices/'+pair['id']+'/revoke',json={},headers=self.headers,environ_overrides=owner).status_code,200)
            self.assertEqual(external.get('/api/state',headers=cookie,environ_overrides=outside).status_code,401)
            self.assertEqual(external.get('/api/state',headers=cookie,environ_overrides=owner).status_code,200)
        finally:os.environ.pop('HA_INGRESS',None)

    def test_ledger_uses_the_shared_url_parser_and_ali_input(self):
        response=self.client.get('/static/shared-link.js')
        try:
            self.assertEqual(response.status_code,200);self.assertIn('extractSharedUrl',response.get_data(as_text=True))
        finally:response.close()
        response=self.client.get('/?camp_hub=1')
        try:
            html=response.get_data(as_text=True);self.assertIn('static/shared-link.js',html);self.assertIn('static/ali-share.js',html);self.assertIn('id="ali-share-text"',html)
        finally:response.close()

if __name__=='__main__':unittest.main()
