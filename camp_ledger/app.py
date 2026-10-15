import io, json, os, re, sqlite3, uuid, zipfile, time, math
from datetime import date, datetime
from urllib.parse import urlsplit
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from flask import Flask, request, jsonify, send_file, abort
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter
from PIL import Image
from live_sync import LiveSync
from exchange import get_rate, dollar_amount, won
from mobile_access import install_mobile, channel_app

app = Flask(__name__, static_folder=str(Path(__file__).with_name('static')))
app.config.update(MAX_FORM_PARTS=None, MAX_FORM_MEMORY_SIZE=None)
DATA = Path(os.environ.get('CAMP_DATA', '/data/camp-ledger'))
DATA.mkdir(parents=True, exist_ok=True)
DB = DATA / 'ledger.sqlite3'
DEFAULT = {'categories':['부품','텐트','의자','테이블','소모품','기타'], 'methods':['온라인','오프라인'], 'buyers':['나','가족'], 'display': {'rowHeight':72,'headerHeight':48,'headerFont':12,'headerPadding':10,'cellPadding':20,'detailFont':15,'titleFont':32,'columns':{'category':{'font':14,'width':130},'name':{'font':16,'width':360},'purchase_price':{'font':16,'width':165},'sale_price':{'font':16,'width':165}}}}
FIELDS = {'category':'구분','name':'상품명','purchase_date':'구입일자','method':'구입처','buyer':'구매자','purchase_price':'구입가격','sale_price':'판매가격','status':'상태','sale_date':'판매일자','disposal_date':'폐기일자','notes':'메모'}
FIELDS.update(purchase_qty='구매수량',current_qty='현재수량',sold_qty='판매수량')
FIELDS.update(cost_unconfirmed='개별원가미확정',shared_purchase_total='공동구매금액')
FIELDS.update(purchase_type='구매유형',brand='브랜드',purchase_url='구매 사이트 링크',shipping_fee='배송비',purchase_currency='최초 구매 통화',foreign_amount='최초 구매 달러금액',exchange_rate='적용 환율',exchange_date='환율 기준일')
FX_FIELDS=('foreign_amount','exchange_rate')
PRICE_FIELDS=('purchase_price','sale_price','shared_purchase_total','shipping_fee')
QTY_FIELDS=('purchase_qty','current_qty','sold_qty')
STATUSES = {'owned':'보유','sold':'판매','disposed':'폐기'}
DISPLAY_LABELS = {**FIELDS, 'purchase_url':'사이트', 'total_purchase_cost':'배송비 포함 구입금액', 'photos':'사진','disposed_qty':'폐기·소진수량','avg_unit_price':'평균단가','current_value':'현재금액','deducted_value':'차감금액','owned_status':'보유 여부','sold_status':'판매 완료','disposed_status':'폐기·소진 여부','product_photo_count':'상품 사진 수','receipt_photo_count':'영수증 사진 수','created_at':'등록일시','updated_at':'수정일시'}
for key in DISPLAY_LABELS:
    DEFAULT['display']['columns'].setdefault(key, {'font':14,'width':320 if key=='notes' else 160})
DEFAULT['display'].update(visibleColumns=['category','name','purchase_price','sale_price'],
                          visibleSummaries=['purchase','sale','net','owned'],photoPickerSource='files',scrollTopPosition='bottom-right',showProductThumb=True,mobileColumns=2,rememberDialogScroll=False,cellVerticalPadding=12,bodyWeight=400,headerWeight=500,titleWeight=700,detailWeight=400)

DEFAULT['display']['mobile']={'font':13,'headerFont':10,'bodyWeight':400,'headerWeight':400,'cellPadding':0,'cellVerticalPadding':0,'rowHeight':0,'headerHeight':0,'headerPadding':0,'cardPadding':12,'fieldGap':12,'columnGap':12,'labelGap':3,'linkSize':30,'columns':{}}

def normalize_settings(s):
    s.setdefault('purchase_types',['미기재','새상품','중고'])
    d=s.setdefault('display',{})
    for key in ('visibleColumns','visibleSummaries','photoPickerSource','scrollTopPosition','showProductThumb','mobileColumns','rememberDialogScroll','headerFont','headerPadding','bodyWeight','headerWeight','titleWeight','detailWeight','cellVerticalPadding'):
        if key not in d: d[key]=list(DEFAULT['display'][key]) if isinstance(DEFAULT['display'][key],list) else DEFAULT['display'][key]
    if 'columnOrder' in d:
        d['columnOrder']=[k for k in d['columnOrder'] if k in DISPLAY_LABELS]+[k for k in DISPLAY_LABELS if k not in d['columnOrder']]
    d.setdefault('addButtonPosition','bottom-right');d.setdefault('addButtonLabel','물품 추가');d.setdefault('fieldLabels',{});d.setdefault('emptyText',{});d.setdefault('photoSize',42);d.setdefault('selectionEnabled',True);d.setdefault('selectionPcEnabled',d['selectionEnabled']);d.setdefault('selectionMobileEnabled',d['selectionEnabled'])
    mobile=d.setdefault('mobile',{})
    if not isinstance(mobile,dict):raise ValueError('모바일 화면 설정을 확인해주세요.')
    for key,value in DEFAULT['display']['mobile'].items():mobile.setdefault(key,{} if key=='columns' else value)
    columns=d.setdefault('columns',{})
    for key,value in DEFAULT['display']['columns'].items():
        if key not in columns: columns[key]=dict(value)
    return s

class InputError(ValueError):
    def __init__(self,message,*fields): super().__init__(message); self.fields=list(fields)

class ManagedConnection(sqlite3.Connection):
    def __exit__(self, *args):
        try: return super().__exit__(*args)
        finally: self.close()

def conn():
    c = sqlite3.connect(DB, timeout=30, factory=ManagedConnection)
    c.row_factory = sqlite3.Row
    c.execute('PRAGMA foreign_keys=ON')
    return c
with conn() as c:
    c.executescript('CREATE TABLE IF NOT EXISTS items(id TEXT PRIMARY KEY, body TEXT NOT NULL); CREATE TABLE IF NOT EXISTS attachments(id TEXT PRIMARY KEY,item_id TEXT NOT NULL REFERENCES items(id) ON DELETE CASCADE,name TEXT NOT NULL,mime TEXT NOT NULL,data BLOB NOT NULL); CREATE TABLE IF NOT EXISTS settings(id INTEGER PRIMARY KEY,body TEXT NOT NULL);')
    if 'kind' not in [r[1] for r in c.execute('PRAGMA table_info(attachments)')]:
        c.execute("ALTER TABLE attachments ADD COLUMN kind TEXT NOT NULL DEFAULT 'product'")
    c.execute('INSERT OR IGNORE INTO settings VALUES(1,?)',(json.dumps(DEFAULT,ensure_ascii=False),))
    c.execute('CREATE TABLE IF NOT EXISTS imports(id TEXT PRIMARY KEY, created REAL NOT NULL, payload TEXT NOT NULL, result TEXT)')

def now(): return datetime.now().astimezone().isoformat(timespec='microseconds')
def settings(c): return normalize_settings(json.loads(c.execute('SELECT body FROM settings WHERE id=1').fetchone()[0]))
def get_item(c, ident):
    r=c.execute('SELECT body FROM items WHERE id=?',(ident,)).fetchone()
    if not r: abort(404, description='상품을 찾을 수 없습니다.')
    item=stock_view(json.loads(r[0]))
    item['attachments']=[dict(x) for x in c.execute('SELECT id,name,mime,kind,length(data) AS size FROM attachments WHERE item_id=? ORDER BY rowid',(ident,))]
    item['product_photo_count']=sum(a['kind']=='product' for a in item['attachments']);item['receipt_photo_count']=sum(a['kind']=='receipt' for a in item['attachments'])
    return item

def fx_meta(raw,day,fetch_rate=True):
    currency=raw.get('purchase_currency') or 'KRW'
    if currency not in ('KRW','USD'): raise InputError('구매 통화를 확인해주세요.','purchase_currency')
    result={'purchase_currency':currency,'foreign_amount':0,'exchange_rate':0,'exchange_date':''}
    if currency=='USD':
        try:
            amount=dollar_amount(raw.get('foreign_amount',0))
            if not day: raise ValueError('달러 환산에는 구매일자가 필요합니다.')
            if fetch_rate:rate=get_rate(day,DB)
            else:
                value=Decimal(str(raw.get('exchange_rate',0)));actual=date.fromisoformat(raw.get('exchange_date',''))
                if not value.is_finite() or not 0<value<1000000 or actual>date.fromisoformat(day):raise ValueError('저장된 환율을 확인해주세요.')
                rate={'exchange_rate':float(value),'exchange_date':actual.isoformat()}
        except ValueError as exc: raise InputError(str(exc),'foreign_amount','purchase_date')
        result.update(foreign_amount=float(amount),exchange_rate=rate['exchange_rate'],exchange_date=rate['exchange_date'])
    return result

@app.get('/api/exchange-rate')
def exchange_rate():
    try: return jsonify(get_rate(request.args.get('date',''),DB))
    except ValueError as exc: raise InputError(str(exc),'purchase_date')

def valid_item(raw,fetch_rate=True):
    obj={k:raw.get(k,'') for k in FIELDS}
    obj['purchase_type']=raw.get('purchase_type') or '미기재'
    obj['purchase_currency']=raw.get('purchase_currency') or 'KRW';obj['exchange_date']=raw.get('exchange_date') or ''
    for k in FIELDS:
        if k not in (*PRICE_FIELDS,*QTY_FIELDS,*FX_FIELDS,'cost_unconfirmed') and not isinstance(obj[k],str): raise ValueError('입력 형식을 확인해주세요.')
    obj['brand']=obj['brand'].strip();obj['purchase_url']=obj['purchase_url'].strip()
    if obj['purchase_url']:
        url=urlsplit(obj['purchase_url'])
        if url.scheme not in ('http','https') or not url.hostname or url.username or url.password: raise InputError('구매 사이트 링크는 http:// 또는 https:// 주소로 입력해주세요.','purchase_url')
    obj['name']=obj['name'].strip()
    if not obj['name']: raise InputError('상품명을 입력해주세요.','name')
    if obj['status'] not in STATUSES: raise ValueError('상태를 확인해주세요.')
    if not obj['purchase_type'].strip() or len(obj['purchase_type'])>80: raise InputError('구매유형을 확인해주세요.','purchase_type')
    obj['cost_unconfirmed']=raw.get('cost_unconfirmed',False)
    if type(obj['cost_unconfirmed']) is not bool: raise ValueError('원가 확정 여부를 확인해주세요.')
    for k in PRICE_FIELDS:
        v=raw.get(k,0)
        if isinstance(v,bool) or not str(v).isdigit(): raise InputError('금액은 0 이상의 정수로 입력해주세요.',k)
        obj[k]=int(v)
        if obj[k]>999999999999: raise InputError('금액이 너무 큽니다.',k)
    for k in ('purchase_date','sale_date','disposal_date','exchange_date'):
        if obj[k]:
            try: date.fromisoformat(obj[k])
            except ValueError: raise InputError('날짜를 확인해주세요.',k)
    for k in QTY_FIELDS:
        default=1 if k=='purchase_qty' else obj['purchase_qty'] if k=='current_qty' and obj['status']=='owned' else obj['purchase_qty'] if k=='sold_qty' and obj['status']=='sold' else 0
        v=raw.get(k,default)
        if v is None or v=='': v=default
        if isinstance(v,bool) or not str(v).isdigit(): raise InputError('수량은 0 이상의 정수로 입력해주세요.',k)
        obj[k]=int(v)
    if obj['purchase_qty']<1 or obj['purchase_qty']>1000000000: raise InputError('구매수량은 1~1,000,000,000개로 입력해주세요.','purchase_qty')
    if obj['status']=='sold' and obj['current_qty']:
        obj['sold_qty']+=obj['current_qty']; obj['current_qty']=0
    elif obj['status']=='disposed':
        obj['current_qty']=0
        if obj['sold_qty']==obj['purchase_qty']: obj['sold_qty']=0
    if obj['current_qty']+obj['sold_qty']>obj['purchase_qty']: raise InputError('현재수량과 판매수량의 합이 구매수량을 넘을 수 없습니다.','purchase_qty','current_qty','sold_qty')
    obj['status']='owned' if obj['current_qty'] else 'sold' if obj['sold_qty'] else 'disposed'
    obj.update(fx_meta(raw,obj['purchase_date'],fetch_rate))
    if obj['purchase_currency']=='USD' and not obj['purchase_price']: obj['purchase_price']=won(obj['foreign_amount'],obj['exchange_rate'])
    return obj

def stock_view(item):
    for key,value in {'purchase_currency':'KRW','foreign_amount':0,'exchange_rate':0,'exchange_date':''}.items():item.setdefault(key,value)
    item.setdefault('brand','');item.setdefault('purchase_url','');item.setdefault('shipping_fee',0)
    item['total_purchase_cost']=item['purchase_price']+item['shipping_fee']
    item.setdefault('purchase_qty',1)
    item.setdefault('current_qty',item['purchase_qty'] if item['status']=='owned' else 0)
    item.setdefault('sold_qty',item['purchase_qty'] if item['status']=='sold' else 0)
    item.setdefault('cost_unconfirmed',False);item.setdefault('shared_purchase_total',0)
    item.setdefault('purchase_type','미기재'); item.setdefault('consumptions',[])
    item['owned_status']='보유' if item['current_qty'] else '미보유'
    item['sold_status']='판매 완료' if item['sold_qty']==item['purchase_qty'] else '일부 판매' if item['sold_qty'] else '미판매'
    item['disposed_status']='폐기·소진' if item['purchase_qty']>item['current_qty']+item['sold_qty'] else '없음'
    item['disposed_qty']=item['purchase_qty']-item['current_qty']-item['sold_qty']
    item['avg_unit_price']=float((Decimal(item['total_purchase_cost'])/item['purchase_qty']).quantize(Decimal('0.01'),rounding=ROUND_HALF_UP))
    item['current_value']=int((Decimal(item['total_purchase_cost'])*item['current_qty']/item['purchase_qty']).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    item['deducted_value']=item['total_purchase_cost']-item['current_value']
    if item['cost_unconfirmed']:
        for key in ('avg_unit_price','current_value','deducted_value'): item[key]=None
    item.setdefault('purchases',[{'id':str(uuid.uuid5(uuid.NAMESPACE_URL,item['id']+'/initial')), 'date':item['purchase_date'],'quantity':item['purchase_qty'],'cost':item['purchase_price'],'shipping_fee':item['shipping_fee'],'notes':'최초 구매'}])
    for purchase in item['purchases']:
        purchase.setdefault('purchase_type',item['purchase_type']);purchase.setdefault('shipping_fee',0)
        for key,value in {'purchase_currency':'KRW','foreign_amount':0,'exchange_rate':0,'exchange_date':''}.items():purchase.setdefault(key,value)
    types={p['purchase_type'] for p in item['purchases']}
    if len(types)>1: item['purchase_type']='혼합'
    elif types: item['purchase_type']=next(iter(types))
    return item

def purchase_ledger(obj,old=None):
    purchases=json.loads(json.dumps(old['purchases'])) if old else []
    if not purchases: purchases=[{'id':str(uuid.uuid4()),'date':obj['purchase_date'],'quantity':obj['purchase_qty'],'cost':obj['purchase_price'],'shipping_fee':obj.get('shipping_fee',0),'notes':'최초 구매','purchase_type':obj['purchase_type']}]
    else:
        quantity=obj['purchase_qty']-sum(p['quantity'] for p in purchases[1:])
        cost=obj['purchase_price']-sum(p['cost'] for p in purchases[1:])
        shipping=obj.get('shipping_fee',0)-sum(p.get('shipping_fee',0) for p in purchases[1:])
        if shipping<0: raise InputError('누적 배송비가 추가 구매 배송비보다 작을 수 없습니다.','shipping_fee')
        if quantity<1 or cost<0: raise InputError('누적 구매수량·금액이 추가 구매 기록보다 작을 수 없습니다.','purchase_qty','purchase_price')
        purchases[0].update(quantity=quantity,cost=cost,shipping_fee=shipping,date=obj['purchase_date'])
        if obj['purchase_type']!='혼합': purchases[0]['purchase_type']=obj['purchase_type']
    purchases[0].update({key:obj.get(key,'KRW' if key=='purchase_currency' else '' if key=='exchange_date' else 0) for key in ('purchase_currency','foreign_amount','exchange_rate','exchange_date')})
    if obj.get('purchase_currency')=='USD':
        purchases[0]['cost']=won(obj['foreign_amount'],obj['exchange_rate']);obj['purchase_price']=sum(p['cost'] for p in purchases)
    return purchases

def edit_records(obj,old,operations,history,stamp):
    if not isinstance(operations,list) or len(operations)>100: raise ValueError('구매·소진 기록을 확인해주세요.')
    purchases=json.loads(json.dumps(old['purchases'])) if old else []
    consumptions=json.loads(json.dumps(old.get('consumptions',[]))) if old else []
    for op in operations:
        kind=op.get('kind')
        if kind in ('restock','consume'):
            q=op.get('quantity'); cost=op.get('cost',0)
            if type(q) is not int or not 1<=q<=1000000000: raise InputError('기록 수량을 확인해주세요.','current_qty')
            if type(cost) is not int or not 0<=cost<=999999999999: raise InputError('구매금액을 확인해주세요.','purchase_price')
            dt=op.get('date',''); notes=op.get('notes','')
            if not isinstance(dt,str) or not isinstance(notes,str): raise ValueError('기록 입력을 확인해주세요.')
            if dt: date.fromisoformat(dt)
            entry={'id':str(uuid.uuid4()),'date':dt,'quantity':q,'notes':notes}
            if kind=='restock':
                pt=op.get('purchase_type') or '미기재'
                if not isinstance(pt,str) or len(pt)>80: raise InputError('구매유형을 확인해주세요.','purchase_type')
                shipping=op.get('shipping_fee',0)
                if type(shipping) is not int or not 0<=shipping<=999999999999: raise InputError('배송비는 0 이상의 정수로 입력해주세요.','shipping_fee')
                meta=op.get('_resolved_meta') or fx_meta(op,dt)
                if meta['purchase_currency']=='USD': cost=won(meta['foreign_amount'],meta['exchange_rate'])
                entry.update(cost=cost,shipping_fee=shipping,purchase_type=pt,**meta); purchases.append(entry)
            else:
                entry['cost']=None if obj['cost_unconfirmed'] else round((obj['purchase_price']+obj['shipping_fee'])*q/obj['purchase_qty']);consumptions.append(entry)
            history.append({'at':stamp,'changes':[{'field':'추가 구매' if kind=='restock' else '사용·소진','before':'','after':f'{q}개'+(f' / {cost:,}원 · 배송비 {shipping:,}원' if kind=='restock' else '')+(' · '+notes if notes else '')}]})
        elif kind in ('undo_purchase','undo_consume'):
            rows=purchases if kind=='undo_purchase' else consumptions
            matches=[(n,p) for n,p in enumerate(rows) if p['id']==op.get('id')]
            if not matches or (kind=='undo_purchase' and matches[0][0]==0): raise ValueError('되돌릴 수 없는 기록입니다.')
            n,entry=matches[0];rows.pop(n)
            history.append({'at':stamp,'changes':[{'field':'구매 되돌리기' if kind=='undo_purchase' else '소진 되돌리기','before':f"{entry['quantity']}개",'after':'취소'}]})
        else: raise ValueError('기록 종류를 확인해주세요.')
    return purchase_ledger(obj,{'purchases':purchases} if purchases else None),consumptions

def valid_settings(s):
    s=normalize_settings(s)
    for key in ('categories','methods','buyers','purchase_types'):
        if not isinstance(s.get(key),list) or any(not isinstance(x,str) or not x.strip() for x in s[key]) or len(set(s[key]))!=len(s[key]): raise ValueError('목록 값을 확인해주세요.')
    d=s['display']
    for key,limit in [('fieldLabels',40),('emptyText',100)]:
        values=d[key]
        if not isinstance(values,dict) or any(k not in DISPLAY_LABELS or (key=='emptyText' and k=='photos') or not isinstance(v,str) or len(v)>limit for k,v in values.items()):raise ValueError('항목 이름·빈 내용 표시 설정을 확인해주세요.')
    if type(d['photoSize']) is not int or not 20<=d['photoSize']<=200:raise ValueError('사진 크기는 20~200px입니다.')
    if any(type(d[k]) is not bool for k in ('selectionPcEnabled','selectionMobileEnabled')):raise ValueError('PC·모바일 선택 기능 설정을 확인해주세요.')
    if type(d['selectionEnabled']) is not bool:raise ValueError('선택 기능 설정을 확인해주세요.')
    limits={'cellVerticalPadding':(0,60),'rowHeight':(24,200),'bodyWeight':(100,900),'headerWeight':(100,900),'titleWeight':(100,900),'detailWeight':(100,900),'headerHeight':(20,120),'headerFont':(10,30),'headerPadding':(0,40),'cellPadding':(0,60),'detailFont':(12,30),'titleFont':(20,48)}
    for k,(low,high) in limits.items():
        if type(d.get(k)) is not int or not low<=d[k]<=high: raise ValueError('화면 크기 설정이 범위를 벗어났습니다.')
    for key in DISPLAY_LABELS:
        col=d['columns'][key]
        if type(col['font']) is not int or not 12<=col['font']<=30 or type(col['width']) is not int or not 24<=col['width']<=1600: raise ValueError('열 설정을 확인해주세요.')
        for prop,(low,high) in {'fontWeight':(100,900),'headerWeight':(100,900),'headerFont':(10,30),'padding':(0,60),'verticalPadding':(0,60),'rowHeight':(24,200),'headerPadding':(0,40),'headerHeight':(20,120)}.items():
            if col.get(prop) is not None and (type(col[prop]) is not int or not low<=col[prop]<=high): raise ValueError('항목별 화면 설정 범위를 확인해주세요.')
        if type(col.get('auto',False)) is not bool: raise ValueError('자동 너비 설정을 확인해주세요.')
    mobile=d['mobile']
    mobile_limits={'font':(10,30),'headerFont':(8,30),'bodyWeight':(100,900),'headerWeight':(100,900),'cellPadding':(0,60),'cellVerticalPadding':(0,60),'rowHeight':(0,200),'headerHeight':(0,120),'headerPadding':(0,40),'cardPadding':(0,40),'fieldGap':(0,60),'columnGap':(0,60),'labelGap':(0,40),'linkSize':(20,64)}
    for key,(low,high) in mobile_limits.items():
        if type(mobile[key]) is not int or not low<=mobile[key]<=high:raise ValueError('모바일 화면 설정 범위를 확인해주세요.')
    if not isinstance(mobile['columns'],dict) or any(k not in DISPLAY_LABELS for k in mobile['columns']):raise ValueError('모바일 항목 설정을 확인해주세요.')
    column_limits={**mobile_limits,'width':(20,600),**{k:(0,60) for k in ('spaceTop','spaceBottom','spaceLeft','spaceRight')},'fontWeight':(100,900),'padding':(0,60),'verticalPadding':(0,60)}
    for col in mobile['columns'].values():
        if not isinstance(col,dict):raise ValueError('모바일 항목 설정을 확인해주세요.')
        for key,value in col.items():
            if key not in column_limits or (value is not None and (type(value) is not int or not column_limits[key][0]<=value<=column_limits[key][1])):raise ValueError('모바일 항목 설정 범위를 확인해주세요.')
    order=d.setdefault('columnOrder',list(DISPLAY_LABELS))
    if not isinstance(order,list) or len(order)!=len(DISPLAY_LABELS) or set(order)!=set(DISPLAY_LABELS): raise ValueError('표시 순서를 확인해주세요.')
    for key,allowed in [('visibleColumns',DISPLAY_LABELS),('visibleSummaries',('purchase','sale','net','owned'))]:
        values=d[key]
        if not isinstance(values,list) or any(not isinstance(v,str) or v not in allowed for v in values) or len(set(values))!=len(values): raise ValueError('메인 화면 표시 항목을 확인해주세요.')
    if not d['visibleColumns']: raise ValueError('목록에 표시할 항목을 하나 이상 선택해주세요.')
    if not isinstance(d['addButtonLabel'],str) or len(d['addButtonLabel'])>40: raise ValueError('추가 버튼 문구는 40자 이하로 입력해주세요.')
    if d['addButtonPosition'] not in ('top-left','top-center','top-right','bottom-left','bottom-center','bottom-right'): raise ValueError('추가 버튼 위치를 확인해주세요.')
    if d['scrollTopPosition'] not in ('top-left','top-center','top-right','bottom-left','bottom-center','bottom-right'): raise ValueError('맨 위로 버튼 위치를 확인해주세요.')
    if d['photoPickerSource'] not in ('gallery','files'): raise ValueError('사진 선택 방식을 확인해주세요.')
    if type(d['mobileColumns']) is not int or not 2<=d['mobileColumns']<=8: raise ValueError('모바일 한 줄 항목 수는 2~8개입니다.')
    if type(d['rememberDialogScroll']) is not bool: raise ValueError('팝업 스크롤 기억 설정을 확인해주세요.')
    if type(d['showProductThumb']) is not bool: raise ValueError('대표 사진 표시 설정을 확인해주세요.')
    return s

RELEASE=os.environ.get('CAMP_VERSION') or re.search(r'version:\s*"([^"]+)"',Path(__file__).with_name('config.yaml').read_text())[1]
sync=LiveSync(app,conn,RELEASE)
mobile=install_mobile(app,conn,settings,DATA)

@app.errorhandler(400)
@app.errorhandler(401)
@app.errorhandler(403)
@app.errorhandler(404)
@app.errorhandler(413)
def http_error(e): return jsonify(error=e.description),e.code
@app.errorhandler(ValueError)
@app.errorhandler(KeyError)
@app.errorhandler(zipfile.BadZipFile)
def value_error(e): return jsonify(error=str(e) or '파일 또는 입력값을 확인해주세요.',fields=getattr(e,'fields',[])),400

@app.before_request
def protect():
    response=mobile['guard']()
    if response is not None: return response
    if request.method in ('POST','PUT','DELETE') and request.headers.get('X-Camp-Request')!='1': abort(400,description='요청을 확인해주세요.')

@app.after_request
def headers(r):
    r.headers['X-Content-Type-Options']='nosniff'
    r.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' blob: data:; connect-src 'self'; object-src 'none'; base-uri 'self'"
    r.headers['Cache-Control']='no-store'
    return r

@app.get('/')
def index(): return app.send_static_file('index.html' if request.args.get('camp_hub')=='1' else 'hub.html')
@app.get('/api/state')
def state():
    with conn() as c:
        c.execute('BEGIN')
        return jsonify(revision=sync.revision(c),release=RELEASE,settings_version=sync.settings_version(c),settings=settings(c),items=[get_item(c,x[0]) for x in c.execute('SELECT id FROM items ORDER BY rowid DESC').fetchall()],mobile=mobile['status']())
@app.put('/api/settings')
def update_settings():
    raw=request.get_json() or {};expected=raw.pop('expected_version',None)
    s=valid_settings(raw)
    with conn() as c:
        c.execute('BEGIN IMMEDIATE')
        if expected is not None and expected!=sync.settings_version(c):return jsonify(error='다른 화면에서 설정이 변경되었습니다. 최신 설정을 불러온 뒤 다시 수정해주세요.'),409
        s['mobile']=settings(c).get('mobile',{})
        c.execute('UPDATE settings SET body=? WHERE id=1',(json.dumps(s,ensure_ascii=False),))
    return jsonify(ok=True)
@app.post('/api/items')
@app.put('/api/items/<ident>')
def save(ident=None):
    raw=json.loads(request.form.get('item','{}'))
    obj=valid_item(raw)
    ops=raw.get('operations',[])
    if not isinstance(ops,list) or len(ops)>100:raise ValueError('구매·소진 기록을 확인해주세요.')
    for op in ops:
        if not isinstance(op,dict):raise ValueError('기록 입력을 확인해주세요.')
        op.pop('_resolved_meta',None)
        if op.get('kind')=='restock':op['_resolved_meta']=fx_meta(op,op.get('date',''))
    delete_ids=json.loads(request.form.get('delete_attachments','[]'))
    if not isinstance(delete_ids,list) or any(not isinstance(x,str) for x in delete_ids): raise ValueError('첨부 목록을 확인해주세요.')
    uploads=[]
    incoming=request.files.getlist('photos')
    kinds=json.loads(request.form.get('photo_kinds','[]'))
    if not kinds: kinds=['product']*len(incoming)
    if len(kinds)!=len(incoming) or any(k not in ('product','receipt') for k in kinds): raise ValueError('사진 종류를 확인해주세요.')
    for f,kind in zip(incoming,kinds):
        data=f.read()
        try:
            with Image.open(io.BytesIO(data)) as im:
                fmt=im.format
                im.verify()
            mime={'JPEG':'image/jpeg','PNG':'image/png','WEBP':'image/webp','GIF':'image/gif'}[fmt]
        except Exception: raise ValueError('사진은 JPG, PNG, WebP, GIF 형식으로 첨부해주세요.')
        uploads.append((str(uuid.uuid4()),Path(f.filename or '사진').name,mime,data,kind))
    with conn() as c:
        c.execute('BEGIN IMMEDIATE')
        old=get_item(c,ident) if ident else None
        if old and raw.get('updated_at')!=old['updated_at']: return jsonify(error='다른 화면에서 변경되었습니다. 최신 기록을 불러온 뒤 다시 수정해주세요.'),409
        ident=ident or str(uuid.uuid4())
        history=old.get('history',[]) if old else []
        stamp=now()
        changes=[{'field':FIELDS[k],'before':old.get(k,''),'after':obj[k]} for k in FIELDS if old and old.get(k)!=obj[k]]
        if not old: changes=[{'field':'등록','before':'','after':obj['name']}]
        removed=[a['name'] for a in old['attachments'] if a['id'] in delete_ids] if old else []
        if uploads: changes.append({'field':'사진 추가','before':'','after':', '.join(u[1] for u in uploads)})
        if removed: changes.append({'field':'사진 삭제','before':', '.join(removed),'after':''})
        if changes: history.append({'at':stamp,'changes':changes})
        obj.update(id=ident,created_at=old['created_at'] if old else stamp,updated_at=stamp,history=history)
        obj['purchases'],obj['consumptions']=edit_records(obj,old,raw.get('operations',[]),history,stamp)
        c.execute('INSERT INTO items VALUES(?,?) ON CONFLICT(id) DO UPDATE SET body=excluded.body',(ident,json.dumps(obj,ensure_ascii=False)))
        for aid in delete_ids: c.execute('DELETE FROM attachments WHERE id=? AND item_id=?',(aid,ident))
        for aid,name,mime,data,kind in uploads: c.execute('INSERT INTO attachments(id,item_id,name,mime,data,kind) VALUES(?,?,?,?,?,?)',(aid,ident,name,mime,data,kind))
    return jsonify(ok=True,id=ident)

@app.post('/api/items/<ident>/restock')
def restock(ident):
    raw=request.get_json() or {}
    quantity=raw.get('quantity'); cost=raw.get('cost'); purchased=raw.get('date','');shipping=raw.get('shipping_fee',0)
    if type(shipping) is not int or not 0<=shipping<=999999999999: raise InputError('배송비는 0 이상의 정수로 입력해주세요.','shipping_fee')
    if type(quantity) is not int or not 1<=quantity<=1000000000: raise ValueError('추가 구매수량을 확인해주세요.')
    if type(cost) is not int or not 0<=cost<=999999999999: raise ValueError('추가 구매금액을 확인해주세요.')
    if purchased: date.fromisoformat(purchased)
    notes=raw.get('notes','')
    purchase_type=raw.get('purchase_type') or '미기재'
    if not isinstance(purchase_type,str) or not purchase_type.strip() or len(purchase_type)>80: raise InputError('구매유형을 확인해주세요.','purchase_type')
    if not isinstance(notes,str): raise ValueError('메모를 확인해주세요.')
    meta=fx_meta(raw,purchased)
    if meta['purchase_currency']=='USD':cost=won(meta['foreign_amount'],meta['exchange_rate'])
    with conn() as c:
        c.execute('BEGIN IMMEDIATE'); item=get_item(c,ident)
        if item['updated_at']!=raw.get('updated_at'): return jsonify(error='다른 화면에서 변경되었습니다. 다시 열어주세요.'),409
        item['purchase_qty']+=quantity;item['current_qty']+=quantity;item['purchase_price']+=cost;item['shipping_fee']+=shipping;item['status']='owned'
        valid_item(item)
        stamp=now();item['updated_at']=stamp
        item['purchases'].append({'id':str(uuid.uuid4()),'date':purchased,'quantity':quantity,'cost':cost,'shipping_fee':shipping,'notes':notes,'purchase_type':purchase_type,**meta})
        item['history'].append({'at':stamp,'changes':[{'field':'추가 구매','before':'','after':f'{quantity}개 / {cost:,}원 · 배송비 {shipping:,}원'+(' · '+notes if notes else '')}]})
        for key in ('attachments','avg_unit_price','current_value','deducted_value','disposed_qty'): item.pop(key,None)
        c.execute('UPDATE items SET body=? WHERE id=?',(json.dumps(item,ensure_ascii=False),ident))
    return jsonify(ok=True)
@app.post('/api/delete')
def delete():
    ids=request.get_json().get('ids',[])
    if not isinstance(ids,list) or any(not isinstance(x,str) for x in ids): raise ValueError('선택 항목을 확인해주세요.')
    with conn() as c:
        c.executemany('DELETE FROM items WHERE id=?',[(x,) for x in ids])
    return jsonify(ok=True)
@app.get('/api/photos/<ident>')
def photo(ident):
    with conn() as c: r=c.execute('SELECT name,mime,data FROM attachments WHERE id=?',(ident,)).fetchone()
    if not r: abort(404)
    return send_file(io.BytesIO(r['data']),mimetype=r['mime'],download_name=r['name'])

def excel_value(v):
    if isinstance(v,str) and v.startswith(('=','+','-','@')): return "'"+v
    return v
@app.get('/api/export')
def export():
    wb=Workbook(); ws=wb.active; ws.title='전체 상품'; hist=wb.create_sheet('변경 이력'); photos=wb.create_sheet('첨부 사진'); prefs=wb.create_sheet('설정'); summary=wb.create_sheet('금액 집계'); purchases_sheet=wb.create_sheet('구매 내역'); consumption_sheet=wb.create_sheet('사용·소진 내역')
    ws.append(['ID']+list(FIELDS.values())+['폐기·소진수량','평균단가','현재금액','차감금액','배송비 포함 구입금액']+['사진 수','사진 파일명','등록 시각','수정 시각'])
    purchases_sheet.append(['상품 ID','상품명','구매일자','구매수량','구입금액','배송비','배송비 포함 금액','개당가격(배송비 포함)','구매유형','메모','구매 통화','달러 구매금액','적용 환율','환율 기준일'])
    consumption_sheet.append(['상품 ID','상품명','사용일자','소진수량','소진금액','메모'])
    hist.append(['상품 ID','상품명','변경 시각','항목','변경 전','변경 후'])
    photos.append(['상품 ID','상품명','첨부 ID','파일명','형식','종류','사진'])
    with conn() as c:
        c.execute('BEGIN')
        items=[get_item(c,r[0]) for r in c.execute('SELECT id FROM items ORDER BY rowid').fetchall()]
        for it in items:
            vals=[STATUSES[it[k]] if k=='status' else ('미확정' if it[k] else '확정') if k=='cost_unconfirmed' else it[k] for k in FIELDS]
            ws.append([excel_value(v) for v in [it['id']]+vals+[it['disposed_qty'],it['avg_unit_price'] if it['avg_unit_price'] is not None else '미확정',it['current_value'] if it['current_value'] is not None else '미확정',it['deducted_value'] if it['deducted_value'] is not None else '미확정',it['total_purchase_cost']]+[len(it['attachments']),', '.join(a['name'] for a in it['attachments']),it['created_at'],it['updated_at']]])
            for consumed in it['consumptions']:
                consumption_sheet.append([excel_value(v) for v in [it['id'],it['name'],consumed['date'],consumed['quantity'],consumed.get('cost') if consumed.get('cost') is not None else '미확정',consumed['notes']]])
            for purchase in it['purchases']:
                purchases_sheet.append([excel_value(v) for v in [it['id'],it['name'],purchase['date'],purchase['quantity'],purchase['cost'],purchase.get('shipping_fee',0),purchase['cost']+purchase.get('shipping_fee',0),round((purchase['cost']+purchase.get('shipping_fee',0))/purchase['quantity'],2) if not it['cost_unconfirmed'] else '미확정',purchase['purchase_type'],purchase['notes'],purchase.get('purchase_currency','KRW'),purchase.get('foreign_amount',0),purchase.get('exchange_rate',0),purchase.get('exchange_date','')]])
            for event in it['history']:
                for ch in event['changes']: hist.append([excel_value(v) for v in [it['id'],it['name'],event['at'],ch['field'],ch['before'],ch['after']]])
            for a in c.execute('SELECT * FROM attachments WHERE item_id=?',(it['id'],)):
                photos.append([excel_value(v) for v in [it['id'],it['name'],a['id'],a['name'],a['mime'],'영수증' if a['kind']=='receipt' else '상품 사진']])
                with Image.open(io.BytesIO(a['data'])) as im:
                    im.thumbnail((180,130)); buf=io.BytesIO(); im.convert('RGB').save(buf,'PNG'); buf.seek(0)
                xl=XLImage(buf); photos.add_image(xl,f'G{photos.max_row}'); photos.row_dimensions[photos.max_row].height=105
        prefs.append(['항목','값']); s=settings(c)
        for k in ('categories','methods','buyers','purchase_types'): prefs.append([{'categories':'구분','methods':'구입처','buyers':'구매자','purchase_types':'구매유형'}[k],json.dumps(s[k],ensure_ascii=False)])
        prefs.append(['화면 설정',json.dumps(s['display'],ensure_ascii=False)])
    summary.append(['상태','수량','구입 합계','실제 판매 합계','순지출'])
    for status,label in list(STATUSES.items())+[('all','전체')]:
        group=[it for it in items if status=='all' or (it['current_qty']>0 if status=='owned' else it['sold_qty']>0 if status=='sold' else it['disposed_qty']>0)]
        p=sum(it['total_purchase_cost'] for it in group); sale=sum(it['sale_price'] for it in group if it['sold_qty']>0)
        summary.append([label,len(group),p,sale,p-sale])
    summary.append(['현재 보유 원가(확정분)',None,sum(it['current_value'] or 0 for it in items),None,None])
    summary.append(['개별 원가 미확정 품목',sum(it['cost_unconfirmed'] for it in items),None,None,None])
    for sheet in wb:
        for cell in sheet[1]:
            if '일자' in str(cell.value) or cell.value=='환율 기준일':
                for row in sheet.iter_rows(min_row=2,min_col=cell.column,max_col=cell.column):
                    value=row[0].value
                    if isinstance(value,str):
                        try:row[0].value=datetime.fromisoformat(value);row[0].number_format='yy.mm.dd'
                        except ValueError:pass
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]: cell.fill=PatternFill('solid',fgColor='214F43'); cell.font=Font(color='FFFFFF',bold=True); cell.alignment=Alignment(vertical='center')
        sheet.row_dimensions[1].height=28
        for col in sheet.columns:
            key=col[0].column_letter; sheet.column_dimensions[key].width=min(55,max(15,max(len(str(cell.value or '')) for cell in col)+3))
        for row in sheet.iter_rows(min_row=2):
            for cell in row: cell.alignment=Alignment(vertical='center',wrap_text=True)
    for n in (6,7):
        for row in ws.iter_rows(min_row=2): row[n].number_format='#,##0"원"'
    photos.column_dimensions['G'].width=28
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf,as_attachment=True,download_name=f'캠핑용품_전체_{date.today()}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.get('/api/import/template')
def import_template():
    with conn() as c: s=settings(c)
    wb=Workbook(); ws=wb.active; ws.title='용품 입력'
    ws.append(list(FIELDS.values()))
    ws.freeze_panes='C2'; ws.auto_filter.ref=f'A1:{get_column_letter(len(FIELDS))}101'; ws.sheet_view.showGridLines=False
    widths=[16,32,17,16,16,18,18,15,17,17,44]+[16]*(len(FIELDS)-11)
    for col,width in zip([get_column_letter(i) for i in range(1,len(FIELDS)+1)],widths): ws.column_dimensions[col].width=width
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='285345'); cell.font=Font(color='FFFFFF',bold=True,size=12)
        cell.alignment=Alignment(vertical='center'); cell.comment=Comment('필수 입력' if cell.value=='상품명' else '선택 입력. 날짜: YYYY-MM-DD / 금액: 원 단위 정수','CAMP LEDGER')
    ws.row_dimensions[1].height=32
    for row in ws.iter_rows(min_row=2,max_row=101,max_col=len(FIELDS)):
        for cell in row:
            cell.font=Font(color='285345',size=11)
            cell.fill=PatternFill('solid',fgColor='F3F6EE' if cell.row%2==0 else 'FFFFFF')
            cell.alignment=Alignment(vertical='center',wrap_text=True)
            if cell.column in (3,9,10): cell.number_format='yyyy-mm-dd'
            elif cell.column in (6,7,12,13,14,16): cell.number_format='#,##0'
            else: cell.number_format='@'
        ws.row_dimensions[row[0].row].height=27
    lists=wb.create_sheet('목록'); lists.sheet_state='hidden'
    for col,key,named,target in [('A','categories','CampCategories','A'),('B','methods','CampMethods','D'),('C','buyers','CampBuyers','E'),('D',None,'CampStatus','H')]:
        values=s[key] if key else ['보유','판매','폐기']
        if not values: continue
        lists[f'{col}1']=key or 'status'
        for n,value in enumerate(values,2):
            lists[f'{col}{n}']=value; lists[f'{col}{n}'].data_type='s'
        wb.defined_names.add(DefinedName(named,attr_text=f"'목록'!${col}$2:${col}${len(values)+1}"))
        dv=DataValidation(type='list',formula1='='+named,allow_blank=True)
        dv.showErrorMessage=not bool(key); dv.errorTitle='상태를 확인해주세요'; dv.error='보유, 판매, 폐기 중에서 선택해주세요.'
        ws.add_data_validation(dv); dv.add(f'{target}2:{target}20001')
    type_column=get_column_letter(list(FIELDS).index('purchase_type')+1)
    dv=DataValidation(type='list',formula1='"'+','.join(s['purchase_types'])+'"',allow_blank=True);ws.add_data_validation(dv);dv.add(f'{type_column}2:{type_column}20001')
    for key in QTY_FIELDS:
        col=get_column_letter(list(FIELDS).index(key)+1)
        dv=DataValidation(type='whole',operator='between',formula1=1 if key=='purchase_qty' else 0,formula2=1000000000,allow_blank=True)
        dv.showErrorMessage=True;dv.error='수량은 정수로 입력해주세요.';ws.add_data_validation(dv);dv.add(f'{col}2:{col}20001')
    for col in [get_column_letter(list(FIELDS).index(key)+1) for key in PRICE_FIELDS]:
        dv=DataValidation(type='whole',operator='between',formula1=0,formula2=999999999999,allow_blank=True)
        dv.showErrorMessage=True; dv.error='금액은 0 이상의 원 단위 정수로 입력해주세요.'
        ws.add_data_validation(dv); dv.add(f'{col}2:{col}20001')
    guide=wb.create_sheet('작성 안내',1); guide.sheet_view.showGridLines=False
    guide.column_dimensions['A'].width=23; guide.column_dimensions['B'].width=95
    instructions=[('캠핑용품 엑셀 업로드','작성한 상품은 기존 장부에 새 항목으로 추가됩니다.'),
        ('작성 위치','용품 입력 시트의 2행부터 상품을 한 줄씩 작성하세요. 제목 행은 유지해주세요.'),
        ('필수 입력','상품명만 필수입니다. 빈 행은 건너뛰고, 예시는 업로드하지 않습니다.'),
        ('금액','원 단위 정수. 비워두면 0원으로 추가합니다. 예: 135000'),
        ('날짜','YYYY-MM-DD 또는 Excel 날짜. 예: 2026-10-13'),
        ('구매유형','새상품 / 중고 / 미기재. 자료에 없으면 미기재로 저장합니다.'),
        ('상태','보유 / 판매 / 폐기. 비워두면 보유로 추가됩니다.'),
        ('구분·구입처·구매자','설정에 등록된 목록을 선택할 수 있습니다. 새 값을 직접 입력하면 목록에도 추가됩니다.'),
        ('업로드 순서','엑셀 업로드 → 파일 선택 → 내용 확인 → 추가하기'),
        ('입력 오류','오류가 있는 행을 수정한 후 다시 올려주세요. 오류가 있으면 어떤 상품도 추가하지 않습니다.'),
        ('반복 업로드','같은 파일을 다시 올리면 같은 상품도 새 항목으로 추가됩니다.'),
        ('수량','구매수량은 누적 구매한 전체 개수, 현재수량은 남은 개수, 판매수량은 판매한 개수입니다. 나머지는 폐기·소진수량입니다.'),
        ('원가 미확정','공동구매로 개별 원가를 모르면 개별원가미확정에 예, 공동구매금액에 총액을 기입합니다. 구입가격은 한 품목에만 합산하고 다른 품목에는 0을 입력해 중복 합산하지 않습니다.'),
        ('사진','이 양식은 상품 정보를 입력합니다. 상품 사진·영수증은 업로드 후 상세 화면에서 첨부하세요.'),
        ('파일 형식','xlsx · 파일당 20,000개 상품 / 20MB까지'),
        ('예시','아래 예시는 작성 안내용이며 업로드 대상이 아닙니다.')]
    for row in instructions: guide.append(row)
    guide.append(['구분: 의자','상품명: 헬리녹스 체어원 / 구입일자: 2026-10-13 / 구입처: 온라인 / 구매자: 나 / 구입가격: 135000 / 상태: 보유'])
    for row in guide:
        guide.row_dimensions[row[0].row].height=36
        for cell in row: cell.alignment=Alignment(vertical='center',wrap_text=True); cell.font=Font(color='52654F',size=11)
        row[0].font=Font(color='285345',bold=True,size=11)
    for cell in guide[1]: cell.fill=PatternFill('solid',fgColor='285345'); cell.font=Font(color='FFFFFF',bold=True,size=12)
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf,as_attachment=True,download_name='캠핑용품_업로드양식.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def import_cell(cell,key):
    value=cell.value
    if value is None or isinstance(value,str) and not value.strip():
        return 0 if key in PRICE_FIELDS else 1 if key=='purchase_qty' else None if key in ('current_qty','sold_qty') else False if key=='cost_unconfirmed' else '미기재' if key=='purchase_type' else 'owned' if key=='status' else ''
    if cell.data_type in ('f','e'): raise ValueError('수식이나 Excel 오류 대신 값을 입력해주세요.')
    if key in FX_FIELDS:
        try:return float(dollar_amount(str(value).replace('$','').replace(',',''))) if key=='foreign_amount' else float(value)
        except (ValueError,TypeError):raise ValueError('달러금액·환율은 숫자로 입력해주세요.')
    if key in (*PRICE_FIELDS,*QTY_FIELDS):
        if isinstance(value,bool): raise ValueError('금액은 0 이상의 정수로 입력해주세요.')
        if isinstance(value,str): value=value.strip().replace(',','').replace('₩','').removesuffix('원').strip()
        if isinstance(value,float):
            if not math.isfinite(value) or not value.is_integer(): raise ValueError('금액에 소수를 사용할 수 없습니다.')
            value=int(value)
        if not str(value).isdigit() or int(value)>999999999999: raise ValueError('금액은 0~999,999,999,999원의 정수로 입력해주세요.')
        return int(value)
    if key in ('purchase_date','sale_date','disposal_date','exchange_date'):
        if isinstance(value,datetime): return value.date().isoformat()
        if isinstance(value,date): return value.isoformat()
        if isinstance(value,str):
            text=value.strip()
            for fmt in ('%Y-%m-%d','%Y/%m/%d','%Y.%m.%d','%Y%m%d','%y.%m.%d'):
                try: return datetime.strptime(text,fmt).date().isoformat()
                except ValueError: pass
        raise ValueError('날짜는 YYYY-MM-DD 또는 Excel 날짜로 입력해주세요.')
    if key=='cost_unconfirmed':
        text=str(value).strip().lower()
        if text in ('예','미확정','true','1'): return True
        if text in ('아니오','확정','false','0'): return False
        raise ValueError('예/아니오 또는 미확정/확정으로 입력해주세요.')
    if key=='status':
        aliases={'보유':'owned','판매':'sold','판매완료':'sold','폐기':'disposed',**{k:k for k in STATUSES}}
        status=aliases.get(str(value).strip().replace(' ',''))
        if not status: raise ValueError('상태는 보유, 판매, 폐기 중에서 입력해주세요.')
        return status
    if isinstance(value,(str,int,float)): return str(value).strip()
    raise ValueError('텍스트로 입력해주세요.')

@app.post('/api/import/preview')
def preview_import():
    f=request.files.get('excel')
    if not f or not (f.filename or '').lower().endswith('.xlsx'): raise ValueError('다운로드한 .xlsx 양식을 선택해주세요.')
    content=f.read(20*1024*1024+1)
    if len(content)>20*1024*1024: raise ValueError('엑셀 파일은 20MB 이하로 올려주세요.')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(x.file_size for x in z.infolist())>120*1024*1024: raise ValueError('엑셀 파일이 너무 큽니다. 나눠서 업로드해주세요.')
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
    except (zipfile.BadZipFile,InvalidFileException,KeyError): raise ValueError('읽을 수 없는 엑셀 파일입니다. .xlsx 파일을 확인해주세요.')
    try:
        ws=next((wb[name] for name in ('용품 입력','전체 상품') if name in wb.sheetnames),wb.worksheets[0])
        if (ws.max_column or 0)>100 or (ws.max_row or 0)>20001: raise ValueError('파일당 최대 20,000행입니다. 입력 범위를 확인해주세요.')
        first=next(ws.iter_rows(min_row=1,max_row=1),())
        mapping={}; labels={**{v:k for k,v in FIELDS.items()},'구매방식':'method','구매 방식':'method'}
        for n,cell in enumerate(first):
            header=str(cell.value or '').strip().removesuffix('*').strip()
            if header not in labels: continue
            key=labels[header]
            if key in mapping: raise ValueError(f'{header} 열이 중복되었습니다.')
            mapping[key]=n
        if 'name' not in mapping: raise ValueError('첫 행에 상품명 열이 필요합니다. 업로드 양식을 내려받아 작성해주세요.')
        rows=[]; errors=[]
        for number,cells in enumerate(ws.iter_rows(min_row=2),2):
            if not any(cells[n].value is not None and str(cells[n].value).strip() for n in mapping.values()): continue
            obj={key:0 if key in PRICE_FIELDS else 1 if key=='purchase_qty' else None if key in ('current_qty','sold_qty') else False if key=='cost_unconfirmed' else '미기재' if key=='purchase_type' else 'owned' if key=='status' else '' for key in FIELDS}
            row_errors=[]
            for key,n in mapping.items():
                try: obj[key]=import_cell(cells[n],key)
                except ValueError as exc: row_errors.append({'row':number,'field':FIELDS[key],'message':str(exc)})
            if not obj['name']: row_errors.append({'row':number,'field':'상품명','message':'상품명을 입력해주세요.'})
            errors.extend(row_errors)
            if not row_errors:
                try: rows.append({'row':number,'item':valid_item(obj)})
                except ValueError as exc: errors.append({'row':number,'field':'수량·입력값','message':str(exc)})
        if not rows and not errors: raise ValueError('추가할 상품이 없습니다. 양식의 2행부터 내용을 작성해주세요.')
        result={'ok':not bool(errors),'count':len(rows),'errors':errors[:100],'error_count':len(errors),'preview':rows[:50],
                'purchase_total':sum(r['item']['purchase_price']+r['item']['shipping_fee'] for r in rows),
                'sale_total':sum(r['item']['sale_price'] for r in rows if r['item']['sold_qty']>0)}
        if not errors:
            token=str(uuid.uuid4())
            with conn() as c:
                c.execute('DELETE FROM imports WHERE created<?',(time.time()-3600,))
                c.execute('INSERT INTO imports VALUES(?,?,?,NULL)',(token,time.time(),json.dumps({'rows':rows,'filename':Path(f.filename).name},ensure_ascii=False)))
            result['token']=token
        return jsonify(result)
    finally: wb.close()

@app.post('/api/import/commit')
def commit_import():
    token=(request.get_json() or {}).get('token')
    if not isinstance(token,str): raise ValueError('업로드 파일을 먼저 확인해주세요.')
    with conn() as c:
        c.execute('BEGIN IMMEDIATE')
        batch=c.execute('SELECT * FROM imports WHERE id=?',(token,)).fetchone()
        if not batch or batch['created']<time.time()-3600: raise ValueError('미리보기가 만료되었습니다. 엑셀 파일을 다시 선택해주세요.')
        if batch['result']: return jsonify(json.loads(batch['result']))
        payload=json.loads(batch['payload']); ids=[]; stamp=now(); s=settings(c)
        for row in payload['rows']:
            it=valid_item(row['item'],fetch_rate=False); ident=str(uuid.uuid4())
            it['purchases']=purchase_ledger(it)
            it.update(id=ident,created_at=stamp,updated_at=stamp,history=[{'at':stamp,'changes':[{'field':'엑셀 등록','before':'','after':f"{payload['filename']} · {row['row']}행"}]}])
            # Always insert under a new ID; existing records and attachments are untouched.
            c.execute('INSERT INTO items VALUES(?,?)',(ident,json.dumps(it,ensure_ascii=False)))
            ids.append(ident)
            for field,key in [('category','categories'),('method','methods'),('buyer','buyers'),('purchase_type','purchase_types')]:
                if it[field] and it[field] not in s[key]: s[key].append(it[field])
        c.execute('UPDATE settings SET body=? WHERE id=1',(json.dumps(s,ensure_ascii=False),))
        result={'ok':True,'count':len(ids),'ids':ids}
        c.execute('UPDATE imports SET result=?,payload=? WHERE id=?',(json.dumps(result),'{}',token))
    return jsonify(result)

@app.get('/api/backup')
def backup():
    buf=io.BytesIO()
    with conn() as c, zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as z:
        c.execute('BEGIN')
        items=[get_item(c,r[0]) for r in c.execute('SELECT id FROM items').fetchall()]
        z.writestr('ledger.json',json.dumps({'version':1,'settings':settings(c),'items':items},ensure_ascii=False))
        for a in c.execute('SELECT id,data FROM attachments'): z.writestr('photos/'+a['id'],a['data'])
    buf.seek(0)
    return send_file(buf,as_attachment=True,download_name=f'캠핑용품_백업_{date.today()}.zip')
@app.post('/api/restore')
def restore():
    f=request.files.get('backup')
    if not f: raise ValueError('백업 파일을 선택해주세요.')
    rows=[]; ids=set(); aids=set()
    with zipfile.ZipFile(f.stream) as z:
        data=json.loads(z.read('ledger.json'))
        if data.get('version')!=1: raise ValueError('지원하지 않는 백업 버전입니다.')
        s=valid_settings(data['settings'])
        for raw in data['items']:
            it=valid_item(raw,fetch_rate=False)
            ident=str(uuid.UUID(raw['id']))
            if ident in ids: raise ValueError('상품 ID가 중복되었습니다.')
            ids.add(ident)
            for k in ('created_at','updated_at'): datetime.fromisoformat(raw[k]); it[k]=raw[k]
            if not isinstance(raw['history'],list): raise ValueError('이력 형식을 확인해주세요.')
            for h in raw['history']:
                datetime.fromisoformat(h['at'])
                for ch in h['changes']:
                    if not isinstance(ch['field'],str) or not isinstance(ch['before'],(str,int)) or not isinstance(ch['after'],(str,int)): raise ValueError('이력 형식을 확인해주세요.')
            it.update(id=ident,history=raw['history']); it['purchases']=raw.get('purchases') or purchase_ledger(it)
            it['consumptions']=raw.get('consumptions',[])
            if not isinstance(it['consumptions'],list): raise ValueError('소진 내역 형식을 확인해주세요.')
            for consumed in it['consumptions']:
                if type(consumed.get('quantity')) is not int or consumed['quantity']<1 or not isinstance(consumed.get('notes'),str): raise ValueError('소진 내역 형식을 확인해주세요.')
                if consumed.get('date'): date.fromisoformat(consumed['date'])
            for purchase in it['purchases']:
                if type(purchase.get('quantity')) is not int or purchase['quantity']<1 or type(purchase.get('cost')) is not int or purchase['cost']<0 or not isinstance(purchase.get('notes'),str): raise ValueError('구매 내역 형식을 확인해주세요.')
                if purchase.get('date'): date.fromisoformat(purchase['date'])
                purchase.setdefault('purchase_type','미기재');purchase.setdefault('shipping_fee',0)
                if type(purchase['shipping_fee']) is not int or not 0<=purchase['shipping_fee']<=999999999999: raise ValueError('배송비 형식을 확인해주세요.')
            if sum(p['quantity'] for p in it['purchases'])!=it['purchase_qty'] or sum(p['cost'] for p in it['purchases'])!=it['purchase_price'] or sum(p['shipping_fee'] for p in it['purchases'])!=it['shipping_fee']: raise ValueError('구매 내역과 누적 수량·금액이 일치하지 않습니다.')
            attachments=[]
            for a in raw['attachments']:
                aid=str(uuid.UUID(a['id']))
                if aid in aids: raise ValueError('사진 ID가 중복되었습니다.')
                aids.add(aid); b=z.read('photos/'+aid)
                with Image.open(io.BytesIO(b)) as im:
                    mime={'JPEG':'image/jpeg','PNG':'image/png','WEBP':'image/webp','GIF':'image/gif'}.get(im.format); im.verify()
                if not mime: raise ValueError('지원하지 않는 사진 형식입니다.')
                kind=a.get('kind','product')
                if kind not in ('product','receipt'): raise ValueError('사진 종류를 확인해주세요.')
                attachments.append((aid,ident,Path(a['name']).name,mime,b,kind))
            rows.append((it,attachments))
    with conn() as c:
        s['mobile']=settings(c).get('mobile',{})
        c.execute('DELETE FROM items')
        c.execute('UPDATE settings SET body=? WHERE id=1',(json.dumps(s,ensure_ascii=False),))
        for it,attachments in rows:
            c.execute('INSERT INTO items VALUES(?,?)',(it['id'],json.dumps(it,ensure_ascii=False)))
            c.executemany('INSERT INTO attachments(id,item_id,name,mime,data,kind) VALUES(?,?,?,?,?,?)',attachments)
    return jsonify(ok=True)

from integrated import install_integration,combined_app
install_integration(app,mobile)

if __name__=='__main__':
    from waitress import serve
    import threading
    threading.Thread(target=lambda:serve(combined_app(app,'external'),host='0.0.0.0',port=8098,threads=16,max_request_body_size=2**40),daemon=True).start()
    serve(combined_app(app,'ingress'),host=os.environ.get('CAMP_HOST','0.0.0.0'),port=int(os.environ.get('PORT','8099')),threads=16,max_request_body_size=2**40)
